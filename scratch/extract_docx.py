"""Dump the F-problem docx to UTF-8 text, inlining OMML math as a linear form."""
import sys
from pathlib import Path
import docx
from docx.oxml.ns import qn

M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'


def omml_to_text(node):
    """Very rough linearisation of an OMML node: keep runs, mark fractions/sub/sup."""
    tag = node.tag
    if tag == M + 'r':
        return ''.join(t.text or '' for t in node.iter(M + 't'))
    if tag == M + 'f':
        num = den = ''
        for ch in node:
            if ch.tag == M + 'num':
                num = ''.join(omml_to_text(c) for c in ch)
            elif ch.tag == M + 'den':
                den = ''.join(omml_to_text(c) for c in ch)
        return f'({num})/({den})'
    if tag in (M + 'sSup', M + 'sSub', M + 'sSubSup'):
        base = sub = sup = ''
        for ch in node:
            if ch.tag == M + 'e':
                base = ''.join(omml_to_text(c) for c in ch)
            elif ch.tag == M + 'sub':
                sub = ''.join(omml_to_text(c) for c in ch)
            elif ch.tag == M + 'sup':
                sup = ''.join(omml_to_text(c) for c in ch)
        out = base
        if sub:
            out += f'_{{{sub}}}'
        if sup:
            out += f'^{{{sup}}}'
        return out
    if tag == M + 'nary':
        chr_el = node.find(M + 'naryPr/' + M + 'chr')
        op = chr_el.get(M + 'val') if chr_el is not None else '∑'
        parts = {'sub': '', 'sup': '', 'e': ''}
        for ch in node:
            key = ch.tag[len(M):]
            if key in parts:
                parts[key] = ''.join(omml_to_text(c) for c in ch)
        return f'{op}_{{{parts["sub"]}}}^{{{parts["sup"]}}}({parts["e"]})'
    if tag == M + 'rad':
        deg = body = ''
        for ch in node:
            if ch.tag == M + 'deg':
                deg = ''.join(omml_to_text(c) for c in ch)
            elif ch.tag == M + 'e':
                body = ''.join(omml_to_text(c) for c in ch)
        return f'root[{deg}]({body})' if deg else f'sqrt({body})'
    if tag == M + 'd':  # delimiter
        inner = ''.join(omml_to_text(c) for c in node)
        return f'({inner})'
    return ''.join(omml_to_text(c) for c in node)


def para_text(p):
    """Walk a paragraph's children in document order so math lands in place."""
    out = []
    for child in p._p:
        if child.tag == qn('w:r'):
            out.append(''.join(t.text or '' for t in child.iter(qn('w:t'))))
        elif child.tag in (M + 'oMath', M + 'oMathPara'):
            out.append(' [[' + omml_to_text(child).strip() + ']] ')
        elif child.tag == qn('w:hyperlink'):
            out.append(''.join(t.text or '' for t in child.iter(qn('w:t'))))
    return ''.join(out).replace('\xa0', ' ')


def main(src, dst):
    d = docx.Document(src)
    lines = []
    for p in d.paragraphs:
        t = para_text(p).rstrip()
        if t.strip():
            style = p.style.name if p.style is not None else ''
            prefix = f'[{style}] ' if style.startswith('Heading') else ''
            lines.append(prefix + t)
    lines.append('')
    lines.append(f'=== TABLE COUNT: {len(d.tables)} ===')
    for i, tb in enumerate(d.tables):
        lines.append(f'--- table {i} ({len(tb.rows)}x{len(tb.columns)}) ---')
        for r in tb.rows:
            lines.append(' | '.join(para_text(c.paragraphs[0]) if c.paragraphs else '' for c in r.cells))
    Path(dst).write_text('\n'.join(lines), encoding='utf-8')
    print('wrote', dst, len(lines), 'lines')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
